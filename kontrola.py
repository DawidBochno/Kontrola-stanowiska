#!/usr/bin/env python3
"""Kontrola bezpieczenstwa stanowiska Windows - lista kontrolna OK / UWAGA
z zaleceniami: antywirus, zapora, aktualizacje, BitLocker, Secure Boot, UAC,
lokalni administratorzy, konto Gosc, SMBv1, PowerShell 2.0, ochrona LSA,
Pulpit zdalny, blokada ekranu, LLMNR, LAPS, system bez wsparcia, miejsce na dysku.
Tylko odczyt - nic nie zmienia w systemie.

Wynik: raport HTML (do wydruku) + wiersz w zbiorczym zestawienie.xlsx.

Uruchomienie: python kontrola.py                       (GUI)
              python kontrola.py --bez-okna [folder]   (bez okna, np. skrypt logowania)
              python kontrola.py --selftest            (test)
"""
import base64
import datetime
import html
import json
import os
import re
import subprocess
import sys
import threading

if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))

KOLUMNY = ["Kontrola", "Wynik", "Ocena", "Zalecenie"]
MAX_DNI_AKTUALIZACJI = 35  # miesieczny cykl poprawek + zapas
MAX_DNI_SYGNATUR = 3
MAX_BLOKADA_S = 900  # 15 minut

# Koniec wsparcia Microsoft (lifecycle) - ta sama tabela co w Audyt AD.
# ponytail: tabela reczna - dopisac nowe wersje Windows, gdy wyjda.
SYSTEMY = [
    (r"windows xp", "2014-04-08"), (r"windows vista", "2017-04-11"),
    (r"windows 7\b", "2020-01-14"), (r"windows 8\.1", "2023-01-10"), (r"windows 8\b", "2016-01-12"),
    (r"server 2003", "2015-07-14"), (r"server 2008", "2020-01-14"), (r"server 2012", "2023-10-10"),
    (r"server 2016", "2027-01-12"), (r"server 2019", "2029-01-09"),
    (r"server 2022", "2031-10-14"), (r"server 2025", "2034-10-10"),
]
WIN10_LTSC = {10240: "2025-10-14", 14393: "2026-10-13", 17763: "2029-01-09", 19044: "2027-01-12"}
WIN10 = "2025-10-14"
WIN11 = {
    22000: ("2023-10-10", "2024-10-08"), 22621: ("2024-10-08", "2025-10-14"),
    22631: ("2025-11-11", "2026-11-10"), 26100: ("2026-10-13", "2027-10-12"),
    26200: ("2027-10-12", "2028-10-10"),
}
BITLOCKER = {  # System.Volume.BitLockerProtection (Eksplorator Windows)
    1: ("włączony", "OK"), 2: ("wyłączony", "UWAGA"), 3: ("szyfrowanie w toku", "INFO"),
    4: ("odszyfrowywanie w toku", "UWAGA"), 5: ("wstrzymany", "UWAGA"),
    6: ("włączony", "OK"), 8: ("czeka na aktywację", "UWAGA"),
}

# Zbieranie faktow: PowerShell 5.1 (jest w kazdym Windows 10/11), bez uprawnien
# administratora. Kazdy odczyt w try - brak danych = null, nie blad calosci.
SKRYPT_PS = r"""
$ErrorActionPreference = 'Stop'
$ProgressPreference = 'SilentlyContinue'
[Console]::OutputEncoding = [Text.Encoding]::UTF8
function Proba($b) { try { & $b } catch { $null } }
function Czas($d) { if ($d) { ([datetime]$d).ToString('s') } else { $null } }
function Rejestr($p, $n) { try { (Get-ItemProperty -Path $p -Name $n).$n } catch { $null } }
$os = Proba { Get-CimInstance Win32_OperatingSystem }
$cs = Proba { Get-CimInstance Win32_ComputerSystem }
$id = [Security.Principal.WindowsIdentity]::GetCurrent()
# Get-BitLockerVolume bez uprawnien czeka ~5 s na dysk, zanim odmowi dostepu
$adm = ([Security.Principal.WindowsPrincipal]$id).IsInRole([Security.Principal.WindowsBuiltInRole]::Administrator)
$sh = Proba { New-Object -ComObject Shell.Application }
$dyski = Proba { Get-CimInstance Win32_LogicalDisk -Filter 'DriveType=3' | ForEach-Object {
    [ordered]@{ dysk = $_.DeviceID; wolne = [double]$_.FreeSpace; rozmiar = [double]$_.Size } } }
$pol = 'HKLM:\SOFTWARE\Microsoft\Windows\CurrentVersion\Policies\System'
$wp = 'HKCU:\Software\Policies\Microsoft\Windows\Control Panel\Desktop'
$wu = 'HKCU:\Control Panel\Desktop'
function Wygaszacz($n) { $v = Rejestr $wp $n; if ($null -eq $v) { $v = Rejestr $wu $n }; $v }
$ts = 'HKLM:\SYSTEM\CurrentControlSet\Control\Terminal Server'
$f = [ordered]@{
  komputer = $env:COMPUTERNAME
  uzytkownik = "$env:USERDOMAIN\$env:USERNAME"
  admin = $adm
  w_domenie = [bool]$cs.PartOfDomain
  domena = $cs.Domain
  system = $os.Caption
  kompilacja = $os.BuildNumber
  wersja = Rejestr 'HKLM:\SOFTWARE\Microsoft\Windows NT\CurrentVersion' 'DisplayVersion'
  uruchomiony = Czas $os.LastBootUpTime
  defender = Proba { $m = Get-MpComputerStatus; [ordered]@{ wlaczony = $m.AntivirusEnabled;
    czas_rzeczywisty = $m.RealTimeProtectionEnabled; sygnatury = Czas $m.AntivirusSignatureLastUpdated;
    tryb = "$($m.AMRunningMode)"; tamper = $m.IsTamperProtected } }
  antywirusy = Proba { Get-CimInstance -Namespace root/SecurityCenter2 -ClassName AntiVirusProduct |
    ForEach-Object { [ordered]@{ nazwa = $_.displayName; stan = [int]$_.productState } } }
  zapora = Proba { Get-NetFirewallProfile | ForEach-Object {
    [ordered]@{ profil = "$($_.Name)"; wlaczona = ("$($_.Enabled)" -eq 'True') } } }
  aktualizacja = Czas (Proba { (New-Object -ComObject Microsoft.Update.AutoUpdate).Results.LastInstallationSuccessDate })
  restart = (Test-Path 'HKLM:\SOFTWARE\Microsoft\Windows\CurrentVersion\WindowsUpdate\Auto Update\RebootRequired') -or
            (Test-Path 'HKLM:\SOFTWARE\Microsoft\Windows\CurrentVersion\Component Based Servicing\RebootPending')
  dyski = $dyski
  bitlocker = Proba { $dyski | ForEach-Object { $d = $_.dysk; [ordered]@{ dysk = $d;
    stan = Proba { [int]$sh.NameSpace($d).Self.ExtendedProperty('System.Volume.BitLockerProtection') };
    ochrona = if ($adm) { Proba { "$((Get-BitLockerVolume -MountPoint $d).ProtectionStatus)" } } } } }
  secureboot = Rejestr 'HKLM:\SYSTEM\CurrentControlSet\Control\SecureBoot\State' 'UEFISecureBootEnabled'
  uac = Rejestr $pol 'EnableLUA'
  uac_admin = Rejestr $pol 'ConsentPromptBehaviorAdmin'
  bezczynnosc = Rejestr $pol 'InactivityTimeoutSecs'
  wygaszacz = [ordered]@{ aktywny = Wygaszacz 'ScreenSaveActive'; haslo = Wygaszacz 'ScreenSaverIsSecure'; czas = Wygaszacz 'ScreenSaveTimeOut' }
  admini = Proba { Get-LocalGroupMember -SID S-1-5-32-544 | ForEach-Object {
    [ordered]@{ nazwa = $_.Name; klasa = "$($_.ObjectClass)"; zrodlo = "$($_.PrincipalSource)" } } }
  uzytkownicy = Proba { Get-LocalUser | ForEach-Object { [ordered]@{ nazwa = $_.Name; wlaczone = $_.Enabled;
    rid = [int]($_.SID.Value.Split('-')[-1]); haslo_wymagane = $_.PasswordRequired } } }
  smb1 = Proba { [int](Get-CimInstance Win32_OptionalFeature -Filter "Name='SMB1Protocol'").InstallState }
  psv2 = Proba { [int](Get-CimInstance Win32_OptionalFeature -Filter "Name='MicrosoftWindowsPowerShellV2Root'").InstallState }
  lsa = Rejestr 'HKLM:\SYSTEM\CurrentControlSet\Control\Lsa' 'RunAsPPL'
  rdp_wylaczony = Rejestr $ts 'fDenyTSConnections'
  rdp_nla = Rejestr "$ts\WinStations\RDP-Tcp" 'UserAuthentication'
  llmnr = Rejestr 'HKLM:\SOFTWARE\Policies\Microsoft\Windows NT\DNSClient' 'EnableMulticast'
  laps = (Rejestr 'HKLM:\SOFTWARE\Microsoft\Windows\CurrentVersion\Policies\LAPS' 'BackupDirectory'),
         (Rejestr 'HKLM:\SOFTWARE\Microsoft\Policies\LAPS' 'BackupDirectory')
  laps_stary = Rejestr 'HKLM:\SOFTWARE\Policies\Microsoft Services\AdmPwd' 'AdmPwdEnabled'
}
$f | ConvertTo-Json -Depth 5 -Compress
"""

# ---------------------------------------------------------- zbieranie ----


def zbierz():
    """Uruchamia SKRYPT_PS i zwraca fakty (slownik)."""
    kod = base64.b64encode(SKRYPT_PS.encode("utf-16-le")).decode()
    r = subprocess.run(
        ["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
         "-EncodedCommand", kod],
        capture_output=True, timeout=300, creationflags=0x08000000)  # CREATE_NO_WINDOW
    wyjscie = r.stdout.decode("utf-8", "replace").lstrip("﻿").strip()
    if not wyjscie.startswith("{"):
        raise RuntimeError("PowerShell nie zwrocil danych: %s" % (
            r.stderr.decode("utf-8", "replace").strip()[:300] or wyjscie[:300]))
    return json.loads(wyjscie.splitlines()[-1])

# ------------------------------------------------------------- ocena ----


def lista(v):
    """ConvertTo-Json zwraca obiekt zamiast listy przy 1 elemencie, null przy 0."""
    if v is None:
        return []
    return v if isinstance(v, list) else [v]


def data(s):
    return datetime.datetime.fromisoformat(s) if s else None


def koniec_wsparcia(system, build):
    """Data konca wsparcia systemu (date) albo None, gdy nieznana."""
    s = (system or "").lower()
    try:
        build = int(build)
    except (TypeError, ValueError):
        build = None
    d = None
    if "windows 11" in s:
        if "ltsc" not in s and build in WIN11:
            d = WIN11[build][any(x in s for x in ("enterprise", "education"))]
    elif "windows 10" in s:
        d = WIN10_LTSC.get(build) if ("ltsc" in s or "ltsb" in s) else WIN10
    else:
        d = next((x for rx, x in SYSTEMY if re.search(rx, s)), None)
    return datetime.date.fromisoformat(d) if d else None


def stan_av(stan):
    """productState z SecurityCenter2 -> (wlaczony, aktualny)."""
    return (stan >> 12) & 0xF == 1, (stan >> 4) & 0xF == 0


def analizuj(f, teraz):
    """Fakty -> lista wierszy [kontrola, wynik, ocena, zalecenie]."""
    W = []
    dodaj = lambda k, w, o, z="": W.append([k, w, o, z if o == "UWAGA" else ""])
    dzis = teraz.date()

    # system
    opis = "%s %s (kompilacja %s)" % (f.get("system") or "?", f.get("wersja") or "", f.get("kompilacja"))
    koniec = koniec_wsparcia(f.get("system"), f.get("kompilacja"))
    zal = "Zaktualizuj system do wspieranej wersji albo wymień komputer."
    if koniec and koniec <= dzis:
        dodaj("System operacyjny", "%s – bez wsparcia od %s" % (opis, koniec), "UWAGA", zal)
    elif koniec and (koniec - dzis).days <= 180:
        dodaj("System operacyjny", "%s – koniec wsparcia %s (za %d dni)" % (
            opis, koniec, (koniec - dzis).days), "UWAGA", "Zaplanuj aktualizację do nowszej wersji (Windows Update).")
    else:
        dodaj("System operacyjny", opis + (" – wsparcie do %s" % koniec if koniec else ""), "OK")

    # antywirus
    df = f.get("defender") or {}
    inne = [a for a in lista(f.get("antywirusy")) if "defender" not in (a.get("nazwa") or "").lower()]
    inne_wl = [a for a in inne if stan_av(a.get("stan") or 0)[0]]
    if df.get("wlaczony") and df.get("czas_rzeczywisty") and df.get("tryb", "Normal") != "Passive mode":
        dodaj("Antywirus", "Microsoft Defender – ochrona w czasie rzeczywistym włączona", "OK")
        sygn = data(df.get("sygnatury"))
        if sygn is None or (teraz - sygn).days > MAX_DNI_SYGNATUR:
            dodaj("Definicje antywirusa", "z %s" % (sygn.strftime("%Y-%m-%d") if sygn else "?"), "UWAGA",
                  "Zaktualizuj definicje: Zabezpieczenia Windows → Ochrona przed wirusami i zagrożeniami "
                  "→ Aktualizacje ochrony.")
        else:
            dodaj("Definicje antywirusa", "z %s" % sygn.strftime("%Y-%m-%d"), "OK")
        dodaj("Ochrona przed naruszeniami (Defender)", "włączona" if df.get("tamper") else "wyłączona",
              "OK" if df.get("tamper") else "UWAGA",
              "Włącz: Zabezpieczenia Windows → Ochrona przed wirusami → Zarządzaj ustawieniami "
              "→ Ochrona przed naruszeniami.")
    elif inne_wl:
        a = inne_wl[0]
        dodaj("Antywirus", "%s – włączony" % a["nazwa"], "OK")
        aktualny = stan_av(a["stan"])[1]
        dodaj("Definicje antywirusa", "aktualne" if aktualny else "nieaktualne", "OK" if aktualny else "UWAGA",
              "Zaktualizuj definicje w programie %s." % a["nazwa"])
    elif not df and not f.get("antywirusy"):
        dodaj("Antywirus", "nie udało się odczytać stanu", "INFO")
    else:
        dodaj("Antywirus", "brak aktywnej ochrony w czasie rzeczywistym", "UWAGA",
              "Włącz Microsoft Defender albo zainstaluj i włącz program antywirusowy.")

    # zapora
    zap = lista(f.get("zapora"))
    wyl = [z["profil"] for z in zap if not z.get("wlaczona")]
    if not zap:
        dodaj("Zapora Windows", "nie udało się odczytać", "INFO")
    elif wyl:
        dodaj("Zapora Windows", "wyłączona dla profilu: " + ", ".join(wyl), "UWAGA",
              "Włącz zaporę dla wszystkich profili (wf.msc), chyba że działa zapora innego programu.")
    else:
        dodaj("Zapora Windows", "włączona (" + ", ".join(z["profil"] for z in zap) + ")", "OK")

    # aktualizacje
    akt = data(f.get("aktualizacja"))
    if akt is None:
        dodaj("Aktualizacje Windows", "brak danych o ostatniej instalacji", "UWAGA",
              "Sprawdź Windows Update (Ustawienia → Windows Update).")
    else:
        dni = (teraz - akt).days
        dodaj("Aktualizacje Windows", "ostatnia instalacja %s (%d dni temu)" % (akt.strftime("%Y-%m-%d"), dni),
              "UWAGA" if dni > MAX_DNI_AKTUALIZACJI else "OK",
              "Uruchom Windows Update i zainstaluj zaległe aktualizacje.")
    if f.get("restart"):
        dodaj("Restart po aktualizacjach", "wymagany", "UWAGA",
              "Uruchom komputer ponownie, aby dokończyć instalację aktualizacji.")

    # szyfrowanie i rozruch
    for b in lista(f.get("bitlocker")):
        opis_bl, ocena = BITLOCKER.get(b.get("stan"), ("niedostępny (np. Windows Home)", "INFO"))
        if b.get("ochrona") == "On":
            opis_bl, ocena = "włączony", "OK"
        elif b.get("ochrona") == "Off" and ocena == "OK":
            opis_bl, ocena = "zaszyfrowany, ochrona wstrzymana", "UWAGA"
        dodaj("Szyfrowanie BitLocker %s" % b.get("dysk"), opis_bl, ocena,
              "Włącz BitLocker (Panel sterowania → Szyfrowanie dysków funkcją BitLocker) "
              "i zapisz klucz odzyskiwania w AD / Entra ID.")
    sb = f.get("secureboot")
    dodaj("Secure Boot", {1: "włączony", 0: "wyłączony"}.get(sb, "niedostępny (tryb BIOS / starszy komputer)"),
          "OK" if sb == 1 else "UWAGA", "Włącz Secure Boot w ustawieniach UEFI komputera.")

    # konta
    uac, uac_admin = f.get("uac"), f.get("uac_admin")
    if uac == 0:
        dodaj("Kontrola konta użytkownika (UAC)", "wyłączona", "UWAGA", "Włącz UAC (EnableLUA = 1) i uruchom ponownie.")
    elif uac_admin == 0:
        dodaj("Kontrola konta użytkownika (UAC)", "włączona, ale bez pytania administratora", "UWAGA",
              "Ustaw UAC na co najmniej domyślny poziom (Panel sterowania → Konta użytkowników).")
    else:
        dodaj("Kontrola konta użytkownika (UAC)", "włączona", "OK")
    admini = lista(f.get("admini"))
    ja = (f.get("uzytkownik") or "").lower()
    if admini:
        nazwy = ", ".join(a["nazwa"] for a in admini)
        jestem = any(a["nazwa"].lower() == ja for a in admini)
        dodaj("Lokalni administratorzy", nazwy, "UWAGA" if jestem else "OK",
              "Zalogowany użytkownik (%s) ma uprawnienia administratora. Do codziennej pracy używaj "
              "konta bez tych uprawnień." % f.get("uzytkownik"))
    else:
        dodaj("Lokalni administratorzy", "nie udało się odczytać", "INFO")
    uz = lista(f.get("uzytkownicy"))
    po_rid = {u.get("rid"): u for u in uz}
    if 501 in po_rid:
        g = po_rid[501]
        dodaj("Konto Gość", "włączone" if g["wlaczone"] else "wyłączone",
              "UWAGA" if g["wlaczone"] else "OK", "Wyłącz konto Gość (lusrmgr.msc).")
    laps = any(v in (1, 2) for v in lista(f.get("laps"))) or f.get("laps_stary") == 1
    if 500 in po_rid:
        a = po_rid[500]
        if not a["wlaczone"]:
            dodaj("Wbudowane konto Administrator", "wyłączone", "OK")
        else:
            dodaj("Wbudowane konto Administrator", "włączone" + (" (hasło zarządzane przez LAPS)" if laps else ""),
                  "OK" if laps else "UWAGA", "Wyłącz konto albo zarządzaj jego hasłem przez LAPS.")
    bez_hasla = [u["nazwa"] for u in uz if u.get("wlaczone") and u.get("haslo_wymagane") is False]
    if uz:
        dodaj("Włączone konta bez wymaganego hasła", ", ".join(bez_hasla) or "brak",
              "UWAGA" if bez_hasla else "OK", "Ustaw hasło albo wyłącz konto.")
    if f.get("w_domenie"):
        dodaj("LAPS (hasło lokalnego administratora)", "włączony" if laps else "brak zasady LAPS",
              "OK" if laps else "UWAGA", "Obejmij komputer zasadą Windows LAPS (GPO lub Intune).")

    # usługi i protokoły
    smb1 = f.get("smb1")
    dodaj("SMB 1.0", "włączony" if smb1 == 1 else "wyłączony", "UWAGA" if smb1 == 1 else "OK",
          "Wyłącz: Funkcje systemu Windows → Obsługa udostępniania plików SMB 1.0/CIFS.")
    if f.get("psv2") == 1:
        dodaj("Windows PowerShell 2.0", "włączony", "UWAGA",
              "Wyłącz: Funkcje systemu Windows → Windows PowerShell 2.0.")
    lsa = f.get("lsa")
    dodaj("Ochrona LSA (haseł w pamięci)", "włączona" if lsa in (1, 2) else "wyłączona",
          "OK" if lsa in (1, 2) else "UWAGA",
          "Włącz: Zabezpieczenia Windows → Zabezpieczenia urządzenia → Izolacja rdzenia → "
          "Ochrona lokalnego urzędu zabezpieczeń (wymaga restartu).")
    if f.get("rdp_wylaczony") == 0:
        nla = f.get("rdp_nla") == 1
        dodaj("Pulpit zdalny (RDP)", "włączony" + (" z NLA" if nla else " bez NLA"), "INFO" if nla else "UWAGA",
              "Włącz uwierzytelnianie na poziomie sieci (NLA) albo wyłącz Pulpit zdalny.")
    else:
        dodaj("Pulpit zdalny (RDP)", "wyłączony", "OK")
    dodaj("LLMNR (rozpoznawanie nazw multiemisji)", "wyłączony" if f.get("llmnr") == 0 else "włączony",
          "OK" if f.get("llmnr") == 0 else "UWAGA",
          "Wyłącz zasadą GPO: Sieć → Klient DNS → Wyłącz rozpoznawanie nazw multiemisji.")

    # blokada ekranu
    try:
        bez = int(f.get("bezczynnosc") or 0)
        ws = f.get("wygaszacz") or {}
        wyg = int(ws.get("czas") or 0) if str(ws.get("aktywny")) == "1" and str(ws.get("haslo")) == "1" else 0
    except ValueError:
        bez = wyg = 0
    czas = min([x for x in (bez, wyg) if x > 0], default=0)
    dodaj("Automatyczna blokada ekranu", "po %d min" % (czas // 60) if czas else "nie ustawiona",
          "OK" if 0 < czas <= MAX_BLOKADA_S else "UWAGA",
          "Ustaw blokadę po maks. 15 min bezczynności (GPO: Logowanie interakcyjne: limit braku "
          "aktywności komputera albo wygaszacz ekranu z hasłem).")

    # dyski
    for d in lista(f.get("dyski")):
        if d.get("rozmiar"):
            proc = 100 * d["wolne"] / d["rozmiar"]
            dodaj("Wolne miejsce %s" % d["dysk"], "%.0f GB z %.0f GB (%.0f%%)" % (
                d["wolne"] / 1e9, d["rozmiar"] / 1e9, proc), "UWAGA" if proc < 10 else "OK",
                "Zwolnij miejsce – przy jego braku aktualizacje się nie instalują.")
    dodaj("Domena", f.get("domena") if f.get("w_domenie") else "poza domeną (%s)" % (f.get("domena") or "grupa robocza"), "INFO")
    up = data(f.get("uruchomiony"))
    if up:
        dodaj("Ostatnie uruchomienie", "%s (%d dni temu)" % (up.strftime("%Y-%m-%d %H:%M"), (teraz - up).days), "INFO")
    return W

# ------------------------------------------------------------- raporty ----


HTML = """<!doctype html>
<html lang="pl"><head><meta charset="utf-8">
<title>Kontrola stanowiska %(komputer)s</title>
<style>
body{font:14px/1.45 Segoe UI,Arial,sans-serif;color:#1d1d1f;margin:32px auto;max-width:1000px;padding:0 16px}
h1{font-size:22px;margin:0 0 4px}
.meta{color:#555;margin-bottom:16px}
.sum span{display:inline-block;padding:4px 10px;border-radius:6px;margin-right:8px;font-weight:600}
table{border-collapse:collapse;width:100%%;margin-top:16px}
th,td{border-bottom:1px solid #ddd;padding:6px 8px;text-align:left;vertical-align:top}
th{background:#f3f3f3}
.UWAGA{background:#ffe1e1}.OK{background:#e3f4e3}.INFO{background:#eef1f6}
td.o{font-weight:600;white-space:nowrap}
.podpis{margin-top:48px;display:flex;gap:64px}.podpis div{flex:1;border-top:1px solid #999;padding-top:4px;color:#555}
.stopka{margin-top:24px;color:#777;font-size:12px}
@media print{body{margin:0}tr{break-inside:avoid}}
</style></head><body>
<h1>Kontrola bezpieczeństwa stanowiska</h1>
<div class="meta">Komputer: <b>%(komputer)s</b> · Użytkownik: %(uzytkownik)s · Data: %(data)s<br>%(tryb)s</div>
<div class="sum"><span class="UWAGA">UWAGA: %(uwagi)d</span><span class="OK">OK: %(ok)d</span><span class="INFO">INFO: %(info)d</span></div>
<table><tr><th>Kontrola</th><th>Wynik</th><th>Ocena</th><th>Zalecenie</th></tr>
%(wiersze)s
</table>
<div class="podpis"><div>Sprawdził (imię, nazwisko, podpis)</div><div>Data</div></div>
<div class="stopka">Raport programu Kontrola stanowiska (github.com/DawidBochno/Kontrola-stanowiska).
Program tylko odczytuje ustawienia, niczego nie zmienia w systemie.</div>
</body></html>
"""


def zapisz_html(f, wiersze, teraz, sciezka):
    e = html.escape
    tr = "\n".join('<tr class="%s"><td>%s</td><td>%s</td><td class="o">%s</td><td>%s</td></tr>' % (
        w[2], e(w[0]), e(str(w[1])), e(w[2]), e(w[3])) for w in wiersze)
    licz = lambda o: sum(1 for w in wiersze if w[2] == o)
    tryb = ("Sprawdzono z uprawnieniami administratora." if f.get("admin") else
            "Sprawdzono bez uprawnień administratora (stan BitLockera z Eksploratora Windows).")
    with open(sciezka, "w", encoding="utf-8") as fh:
        fh.write(HTML % dict(komputer=e(f.get("komputer") or ""), uzytkownik=e(f.get("uzytkownik") or ""),
                             data=teraz.strftime("%Y-%m-%d %H:%M"), tryb=tryb, wiersze=tr,
                             uwagi=licz("UWAGA"), ok=licz("OK"), info=licz("INFO")))


def zapisz_zestawienie(f, wiersze, teraz, sciezka):
    """Jeden wiersz na komputer (ponowna kontrola zastepuje wiersz),
    kolumna na kontrole - zbiorczy widok np. po obejsciu biura z pendrive."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    kolory = {"UWAGA": PatternFill("solid", fgColor="FFC7CE"), "OK": PatternFill("solid", fgColor="C6EFCE")}
    if os.path.exists(sciezka):
        wb = openpyxl.load_workbook(sciezka)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Zestawienie"
        ws.append(["Komputer", "Data", "Użytkownik", "System", "Liczba UWAG"])
    naglowki = [c.value for c in ws[1]]
    for w in wiersze:
        if w[0] not in naglowki and w[2] != "INFO":
            naglowki.append(w[0])
            ws.cell(1, len(naglowki), w[0])
    nr = next((r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == f.get("komputer")), ws.max_row + 1)
    for c in range(1, len(naglowki) + 1):
        ws.cell(nr, c).value = None
        ws.cell(nr, c).fill = PatternFill()
    for i, v in enumerate([f.get("komputer"), teraz.replace(microsecond=0), f.get("uzytkownik"), f.get("system"),
                           sum(1 for w in wiersze if w[2] == "UWAGA")], 1):
        ws.cell(nr, i, v)
    ws.cell(nr, 2).number_format = "yyyy-mm-dd hh:mm"
    for w in wiersze:
        if w[0] in naglowki and w[2] != "INFO":
            c = ws.cell(nr, naglowki.index(w[0]) + 1, str(w[1]))
            c.fill = kolory.get(w[2], PatternFill())
    for c in ws[1]:
        c.font = Font(bold=True)
        ws.column_dimensions[c.column_letter].width = 22
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(sciezka)


def run(out_dir, log, zrodlo=None):
    """Kontrola, raport HTML, wiersz w zestawieniu. Zwraca (fakty, wiersze, plik_html)."""
    log("Sprawdzanie ustawien komputera ...")
    f = (zrodlo or zbierz)()
    teraz = datetime.datetime.now()
    wiersze = analizuj(f, teraz)
    uwagi = sum(1 for w in wiersze if w[2] == "UWAGA")
    log("%s: %d kontroli, %d z UWAGA" % (f.get("komputer"), len(wiersze), uwagi))
    os.makedirs(out_dir, exist_ok=True)
    plik = os.path.join(out_dir, "kontrola_%s_%s.html" % (f.get("komputer"), teraz.strftime("%Y-%m-%d")))
    zapisz_html(f, wiersze, teraz, plik)
    log("Raport: " + plik)
    zest = os.path.join(out_dir, "zestawienie.xlsx")
    try:
        zapisz_zestawienie(f, wiersze, teraz, zest)
        log("Zestawienie: " + zest)
    except PermissionError:
        log("BLAD: zestawienie.xlsx jest otwarte (np. w Excelu) - nie dopisano wiersza. Raport HTML zapisany.")
    return f, wiersze, plik

# ---------------------------------------------------------------- GUI ----


def gui():
    import ctypes
    import tkinter as tk
    from tkinter import messagebox, scrolledtext, ttk

    root = tk.Tk()
    root.title("Kontrola bezpieczenstwa stanowiska")
    root.geometry("1100x660")
    pad = dict(padx=6, pady=3)
    v_out = tk.StringVar(value=os.path.join(APP_DIR, "OUTPUT"))
    stan = {"plik": None}

    f = ttk.Frame(root)
    f.pack(fill="x", **pad)
    ttk.Label(f, text="Folder wyjsciowy:").pack(side="left", **pad)
    ttk.Entry(f, textvariable=v_out, width=70).pack(side="left", **pad)
    admin = bool(ctypes.windll.shell32.IsUserAnAdmin())
    ttk.Label(root, text="Uprawnienia: administrator" if admin else
              "Uprawnienia: zwykly uzytkownik (wystarcza; jako administrator dokladniejszy stan BitLockera)",
              foreground="gray").pack(anchor="w", padx=12)

    h = ttk.Frame(root)
    h.pack(fill="x", **pad)
    btn = ttk.Button(h, text="Sprawdz")
    btn.pack(side="left", padx=6)
    b_html = ttk.Button(h, text="Otworz raport", state="disabled",
                        command=lambda: os.startfile(stan["plik"]))
    b_html.pack(side="left", padx=6)
    ttk.Button(h, text="Otworz zestawienie", command=lambda: (
        os.startfile(os.path.join(v_out.get(), "zestawienie.xlsx"))
        if os.path.exists(os.path.join(v_out.get(), "zestawienie.xlsx"))
        else log("Zestawienie jeszcze nie istnieje - kliknij Sprawdz."))).pack(side="left", padx=6)

    def jako_admin():
        skrypt = os.path.abspath(__file__)
        if ctypes.windll.shell32.ShellExecuteW(None, "runas", sys.executable, '"%s"' % skrypt, APP_DIR, 1) > 32:
            root.destroy()

    if not admin:
        ttk.Button(h, text="Uruchom jako administrator", command=jako_admin).pack(side="left", padx=6)

    ramka = ttk.Frame(root)
    ramka.pack(fill="both", expand=True, **pad)
    t = ttk.Treeview(ramka, columns=KOLUMNY, show="headings")
    for k, szer in zip(KOLUMNY, (260, 380, 70, 380)):
        t.heading(k, text=k)
        t.column(k, width=szer, anchor="w")
    for o, kolor in (("UWAGA", "#FFD7D7"), ("OK", "#DFF2DF"), ("INFO", "#EEF1F6")):
        t.tag_configure(o, background=kolor)
    t.bind("<Double-1>", lambda ev: t.focus() and messagebox.showinfo(
        "Szczegoly", "\n\n".join("%s: %s" % kv for kv in zip(KOLUMNY, t.item(t.focus(), "values")) if kv[1])))
    sb = ttk.Scrollbar(ramka, orient="vertical", command=t.yview)
    t.configure(yscrollcommand=sb.set)
    sb.pack(side="right", fill="y")
    t.pack(side="left", fill="both", expand=True)

    log_box = scrolledtext.ScrolledText(root, height=5)
    log_box.pack(fill="x", **pad)

    def log(msg):
        def put():
            log_box.insert("end", str(msg) + "\n")
            log_box.see("end")
        root.after(0, put)

    def pokaz(wiersze, plik):
        t.delete(*t.get_children())
        kolejnosc = {"UWAGA": 0, "OK": 1, "INFO": 2}
        for w in sorted(wiersze, key=lambda w: kolejnosc[w[2]]):
            t.insert("", "end", values=w, tags=(w[2],))
        stan["plik"] = plik
        b_html.config(state="normal")

    def start():
        log_box.delete("1.0", "end")
        btn.config(state="disabled")

        def work():
            try:
                _, wiersze, plik = run(v_out.get().strip('" '), log)
                root.after(0, lambda: pokaz(wiersze, plik))
            except Exception as e:
                log("BLAD: %s" % e)
            finally:
                root.after(0, lambda: btn.config(state="normal"))

        threading.Thread(target=work, daemon=True).start()

    btn.config(command=start)
    if "--selftest" in sys.argv:
        root.after(200, root.destroy)
    import aktualizacja
    aktualizacja.start(root, "DawidBochno/Kontrola-stanowiska", "main", "kontrola.py")
    root.mainloop()

# ------------------------------------------------------------- selftest ----


def selftest():
    import tempfile
    import openpyxl

    teraz = datetime.datetime(2026, 10, 7, 12, 0)
    iso = lambda dni: (teraz - datetime.timedelta(days=dni)).strftime("%Y-%m-%dT%H:%M:%S")
    dobry = {
        "komputer": "UG-PC-001", "uzytkownik": "URZAD\\jkowalska", "admin": False,
        "w_domenie": True, "domena": "urzad.local", "system": "Microsoft Windows 11 Pro",
        "kompilacja": "26200", "wersja": "25H2", "uruchomiony": iso(1),
        "defender": {"wlaczony": True, "czas_rzeczywisty": True, "sygnatury": iso(1),
                     "tryb": "Normal", "tamper": True},
        "antywirusy": {"nazwa": "Windows Defender", "stan": 397568},  # 1 element = obiekt, nie lista
        "zapora": [{"profil": p, "wlaczona": True} for p in ("Domain", "Private", "Public")],
        "aktualizacja": iso(5), "restart": False,
        "dyski": [{"dysk": "C:", "wolne": 200e9, "rozmiar": 500e9}],
        "bitlocker": {"dysk": "C:", "stan": 1, "ochrona": None},
        "secureboot": 1, "uac": 1, "uac_admin": 5, "bezczynnosc": 600,
        "wygaszacz": {"aktywny": "1", "haslo": None, "czas": None},
        "admini": [{"nazwa": "UG-PC-001\\Administrator", "klasa": "User", "zrodlo": "Local"},
                   {"nazwa": "URZAD\\Administratorzy domeny", "klasa": "Group", "zrodlo": "ActiveDirectory"}],
        "uzytkownicy": [{"nazwa": "Administrator", "wlaczone": False, "rid": 500, "haslo_wymagane": True},
                        {"nazwa": "Gość", "wlaczone": False, "rid": 501, "haslo_wymagane": False},
                        {"nazwa": "WsiAccount", "wlaczone": False, "rid": 1002, "haslo_wymagane": False}],
        "smb1": 2, "psv2": None, "lsa": 2, "rdp_wylaczony": 1, "rdp_nla": 1, "llmnr": 0,
        "laps": [2, None], "laps_stary": None,
    }
    W = analizuj(dobry, teraz)
    oc = {w[0]: w for w in W}
    assert all(len(w) == 4 for w in W)
    assert not [w for w in W if w[2] == "UWAGA"], [w for w in W if w[2] == "UWAGA"]
    assert all(not w[3] for w in W)  # zalecenia tylko przy UWAGA
    assert oc["System operacyjny"][1].endswith("wsparcie do 2027-10-12")
    assert oc["Automatyczna blokada ekranu"][1] == "po 10 min"
    assert oc["Szyfrowanie BitLocker C:"][1] == "włączony" and oc["LAPS (hasło lokalnego administratora)"][2] == "OK"
    assert "Windows PowerShell 2.0" not in oc

    zly = dict(dobry, system="Microsoft Windows 10 Pro", kompilacja="19045", aktualizacja=iso(80), restart=True,
               defender={"wlaczony": False, "czas_rzeczywisty": False, "sygnatury": iso(30), "tryb": "Normal",
                         "tamper": False},
               antywirusy=[{"nazwa": "Windows Defender", "stan": 393472}],
               zapora=[{"profil": "Domain", "wlaczona": True}, {"profil": "Public", "wlaczona": False}],
               dyski=[{"dysk": "C:", "wolne": 10e9, "rozmiar": 250e9}],
               bitlocker=[{"dysk": "C:", "stan": 2, "ochrona": None}, {"dysk": "D:", "stan": 0, "ochrona": None}],
               secureboot=None, uac=1, uac_admin=0, bezczynnosc=None,
               wygaszacz={"aktywny": "1", "haslo": "1", "czas": "1800"},
               admini=[{"nazwa": "URZAD\\jkowalska", "klasa": "User", "zrodlo": "ActiveDirectory"}],
               uzytkownicy=[{"nazwa": "Administrator", "wlaczone": True, "rid": 500, "haslo_wymagane": True},
                            {"nazwa": "Gość", "wlaczone": True, "rid": 501, "haslo_wymagane": False}],
               smb1=1, psv2=1, lsa=None, rdp_wylaczony=0, rdp_nla=0, llmnr=None, laps=None)
    W = analizuj(zly, teraz)
    oc = {w[0]: w for w in W}
    uw = {w[0] for w in W if w[2] == "UWAGA"}
    assert uw == {
        "System operacyjny", "Antywirus", "Zapora Windows", "Aktualizacje Windows", "Restart po aktualizacjach",
        "Szyfrowanie BitLocker C:", "Secure Boot", "Kontrola konta użytkownika (UAC)", "Lokalni administratorzy",
        "Konto Gość", "Wbudowane konto Administrator", "Włączone konta bez wymaganego hasła",
        "LAPS (hasło lokalnego administratora)", "SMB 1.0", "Windows PowerShell 2.0",
        "Ochrona LSA (haseł w pamięci)", "Pulpit zdalny (RDP)", "LLMNR (rozpoznawanie nazw multiemisji)",
        "Automatyczna blokada ekranu", "Wolne miejsce C:"}, uw
    assert all(w[3] for w in W if w[2] == "UWAGA")
    assert oc["System operacyjny"][1].endswith("bez wsparcia od 2025-10-14")
    assert oc["Zapora Windows"][1] == "wyłączona dla profilu: Public"
    assert oc["Szyfrowanie BitLocker D:"][2] == "INFO" and oc["Automatyczna blokada ekranu"][1] == "po 30 min"
    assert oc["Włączone konta bez wymaganego hasła"][1] == "Gość"
    assert "jkowalska" in oc["Lokalni administratorzy"][3]

    # antywirus innej firmy, Defender pasywny; BitLocker przez Get-BitLockerVolume
    inny = dict(dobry, defender={"wlaczony": True, "czas_rzeczywisty": False, "tryb": "Passive mode"},
                antywirusy=[{"nazwa": "Windows Defender", "stan": 393472}, {"nazwa": "ESET Security", "stan": 266240}],
                bitlocker={"dysk": "C:", "stan": 1, "ochrona": "Off"})
    oc = {w[0]: w for w in analizuj(inny, teraz)}
    assert oc["Antywirus"][1] == "ESET Security – włączony" and oc["Definicje antywirusa"][2] == "OK"
    assert oc["Szyfrowanie BitLocker C:"][1:3] == ["zaszyfrowany, ochrona wstrzymana", "UWAGA"]
    assert stan_av(266240) == (True, True) and stan_av(393472) == (False, True)
    assert koniec_wsparcia("Microsoft Windows 11 Enterprise", "22631") == datetime.date(2026, 11, 10)
    assert koniec_wsparcia("Microsoft Windows Server 2022 Standard", "20348") == datetime.date(2031, 10, 14)
    assert koniec_wsparcia("Microsoft Windows 11 Pro", "29999") is None
    assert lista(None) == [] and lista({"a": 1}) == [{"a": 1}]

    # raporty: HTML + zestawienie (aktualizacja wiersza, nowa kolumna)
    tmp = tempfile.mkdtemp()
    lines = []
    f1, w1, plik = run(tmp, lines.append, zrodlo=lambda: dobry)
    assert all(ord(ch) < 128 for line in lines for ch in line), lines
    tresc = open(plik, encoding="utf-8").read()
    assert "UG-PC-001" in tresc and "UWAGA: 0" in tresc and "&lt;" not in tresc
    run(tmp, lines.append, zrodlo=lambda: dict(zly, komputer="UG-PC-002"))
    run(tmp, lines.append, zrodlo=lambda: dict(dobry, smb1=1))  # ponownie PC-001: zastapic wiersz
    ws = openpyxl.load_workbook(os.path.join(tmp, "zestawienie.xlsx")).active
    assert [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)] == ["UG-PC-001", "UG-PC-002"]
    nag = [c.value for c in ws[1]]
    assert nag[:5] == ["Komputer", "Data", "Użytkownik", "System", "Liczba UWAG"] and "Domena" not in nag
    assert "Windows PowerShell 2.0" in nag and "Szyfrowanie BitLocker D:" not in nag  # INFO bez kolumny
    smb = nag.index("SMB 1.0") + 1
    assert ws.cell(2, smb).value == "włączony" and ws.cell(2, smb).fill.fgColor.rgb.endswith("FFC7CE")
    assert ws.cell(2, 5).value == 1 and ws.cell(3, 5).value == 20

    # prawdziwy odczyt z tego komputera (w CI: czysty Windows Server)
    if sys.platform == "win32":
        f = zbierz()
        for k in ("komputer", "system", "kompilacja", "zapora", "uzytkownicy", "dyski"):
            assert f.get(k), (k, f.get(k))
        W = analizuj(f, datetime.datetime.now())
        assert len(W) > 15 and {"System operacyjny", "Zapora Windows"} <= {w[0] for w in W}
        print("ten komputer: %d kontroli, %d z UWAGA" % (len(W), sum(1 for w in W if w[2] == "UWAGA")))

    import aktualizacja
    aktualizacja.selftest()
    print("selftest OK")


def bez_okna():
    """--bez-okna [folder]: kontrola bez okna (np. skrypt logowania, zadanie)."""
    i = sys.argv.index("--bez-okna")
    out = sys.argv[i + 1] if len(sys.argv) > i + 1 else os.path.join(APP_DIR, "OUTPUT")
    run(out, print)


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
        if "--gui" in sys.argv:
            gui()
    elif "--bez-okna" in sys.argv:
        bez_okna()
    else:
        gui()
